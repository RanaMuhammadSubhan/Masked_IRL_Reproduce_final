"""Build the technical narrative using frozen canonical CSVs and existing figures."""
from pathlib import Path
import csv,json,hashlib,argparse,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ap=argparse.ArgumentParser();ap.add_argument('--package',type=Path,required=True);ap.add_argument('--output',type=Path,required=True);a=ap.parse_args();P=a.package
def rows(name,group='primary'):return list(csv.DictReader((P/'results/tables'/group/name).open(encoding='utf-8-sig')))
S=rows('summary_statistics.csv');E=rows('summary_statistics.csv','extensions');D=rows('paired_method_differences.csv','extensions')
def stat(method,ep,p=0,source='oracle'):
 r=next(r for r in S if r['method']==method and r['endpoint']==ep and float(r['p_flip'])==p and r['mask_type']==source);return f"{float(r['mean']):.2f} ± {float(r['se']):.2f}"
def ext(variant,lam,ep,p=0):
 r=next(r for r in E if r['variant']==variant and float(r['lambda_mask'])==lam and r['endpoint']==ep and float(r['p_flip'])==p);return f"{float(r['mean']):.2f} ± {float(r['se']):.2f}"
def diff(p,ep,comp):
 r=next(r for r in D if r['variant']=='uncertainty' and float(r['p_flip'])==p and r['endpoint']==ep and r['comparison']=='extension minus '+comp);return f"{float(r['mean']):+.2f} ± {float(r['se']):.2f}"
EP=['pre_adapt_heldout','post_adapt_heldout','strict_zero_shot']
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.72);sec.left_margin=sec.right_margin=Inches(.8);sec.header_distance=sec.footer_distance=Inches(.3)
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 st=doc.styles[n];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=1.1;doc.styles['Normal'].paragraph_format.space_after=Pt(7)
for n,size in [('Title',26),('Subtitle',15),('Heading 1',17),('Heading 2',12)]:
 doc.styles[n].font.size=Pt(size);doc.styles[n].paragraph_format.space_after=Pt(10)
doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].font.italic=False
sec.header.paragraphs[0].text='MASKED IRL  |  IMPLEMENTATION REPORT AND RESEARCH LEARNINGS'
sec.header.paragraphs[0].runs[0].font.size=Pt(8)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=fp.add_run('7 October 2026  •  ');r.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)
doc.core_properties.title='Masked IRL Reproduction Implementation and Learnings';doc.core_properties.author='';doc.core_properties.subject='Partially paper-faithful reconstruction on a controlled replacement simulation dataset'
toc=[];page=1;table_no=0;pending_break=False
def br():
 global page,pending_break
 pending_break=True;page+=1
def h(text,level=1,toc_entry=True):
 global pending_break
 p=doc.add_heading(text,level)
 if pending_break:p.paragraph_format.page_break_before=True;pending_break=False
 if level==1 and toc_entry:
  key='sec'+str(len(toc));b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(len(toc)));b.set(qn('w:name'),key);p._p.insert(0,b);end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(len(toc)));p._p.append(end);toc.append((text,page,key))
 return p
def p(s):doc.add_paragraph(s)
def table(caption,headers,data,widths=None):
 global table_no
 table_no+=1;cap=doc.add_paragraph(f'Table {table_no}. {caption}','Caption');cap.paragraph_format.keep_with_next=True
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths is None:widths=[6.9/len(headers)]*len(headers)
 for col,width in zip(t.columns,widths):col.width=Inches(width)
 for i,v in enumerate(headers):t.rows[0].cells[i].text=v
 for row in data:
  for cell,v in zip(t.add_row().cells,row):cell.text=str(v)
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
  if ri==0:pr.append(OxmlElement('w:tblHeader'))
  for j,cell in enumerate(row.cells):
   cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tc=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');mar.append(e)
   tc.append(mar);borders=OxmlElement('w:tcBorders')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tc.append(borders)
   if ri==0:
    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'EEEEEE');tc.append(fill)
   for pp in cell.paragraphs:
    pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(2);pp.paragraph_format.line_spacing=1.03
    for rr in pp.runs:rr.font.size=Pt(9.5);rr.bold=ri==0
 doc.add_paragraph().paragraph_format.space_after=Pt(2)
def figure(name,caption):
 doc.add_picture(str(P/'figures'/name),width=Inches(6.85));doc.paragraphs[-1].paragraph_format.space_after=Pt(3);doc.paragraphs[-1].paragraph_format.keep_with_next=True
 doc.add_paragraph(caption,'Caption')
def eq(s):
 pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
 math=OxmlElement('m:oMathPara');o=OxmlElement('m:oMath')
 def run(text,plain=False):
  r=OxmlElement('m:r')
  if plain:
   pr=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');pr.append(sty);r.append(pr)
  t=OxmlElement('m:t');t.text=text;r.append(t);return r
 known={'ℒIRL':('ℒ','IRL'),'ℒmask':('ℒ','mask'),'cdemo':('c','demo'),'ccontrast':('c','contrast'),'Fdemo,j':('F','demo,j'),'Ftrain,j':('F','train,j'),'Eε':('E','ε')}
 pattern=r'ℒIRL|ℒmask|cdemo|ccontrast|Fdemo,j|Ftrain,j|Eε|[A-Za-zΣℓ][ᵢⱼₜₖₓ]+|SD²'
 last=0
 for match in re.finditer(pattern,s):
  if match.start()>last:o.append(run(s[last:match.start()]))
  token=match.group();base,sub=known.get(token,(token[0],token[1:].translate(str.maketrans('ᵢⱼₜₖₓ','ijtkx'))))
  if token=='SD²':base,sub='SD','2'
  script=OxmlElement('m:sSup' if token=='SD²' else 'm:sSub');e=OxmlElement('m:e');e.append(run(base));snode=OxmlElement('m:sup' if token=='SD²' else 'm:sub');snode.append(run(sub,sub in ['IRL','mask','demo','contrast','demo,j','train,j']));script.append(e);script.append(snode);o.append(script);last=match.end()
 if last<len(s):o.append(run(s[last:]))
 math.append(o);pp._p.append(math)

doc.add_paragraph('Masked IRL: Reproduction, Paper-Faithful Reconstruction, Robustness Analysis, and Exploratory Improvement','Title')
doc.add_paragraph('Implementation Report and Research Learnings','Subtitle')
p('7 October 2026')
p('Source paper: Masked IRL: LLM-Guided Reward Disambiguation from Demonstrations and Language. Manuscript v1, arXiv:2511.14565.')
p('Source repository: MIT-CLEAR-Lab/Masked-IRL\nCommit: b52bc2e3f1ec5597360a74c2641283f6c984a76b')
h('Reproduction classification',2)
p('Partially Paper-Faithful Reconstruction on a Controlled Replacement Simulation Dataset')
p('This report documents completed experiments and their limitations. It accompanies the unchanged short professor report, canonical result tables, runnable code and independent verification evidence. No new training or data generation was performed to prepare this delivery.')
br();h('Executive summary',toc_entry=False)
p('The assignment was to understand an unfamiliar inverse reinforcement learning method, reproduce useful baselines, explain discrepancies and investigate a meaningful limitation. The released learning code was usable, but the original central simulation trajectories and environment assets were unavailable. Exact numerical reproduction of the paper was therefore not possible. The completed work reconstructs a controlled benchmark while preserving a clear distinction between original evidence, implementation choices and exploratory results.')
p('A released-code V2 baseline established that the pipeline could learn and evaluate rewards on replacement trajectories. V3 then restored recoverable manuscript settings: a frozen pinned T5 encoder, local coordinate-wise invariance regularization, Uniform perturbations, lambda 10, and a 1000-epoch pretraining plus 100-epoch adaptation schedule. The released cross-sample importance estimator was deliberately retained and recorded. This is a partially paper-faithful reconstruction, not an exact original simulation or physical-robot reproduction.')
p('Across three prespecified seeds, Explicit Mask had the highest adapted core mean, 70.38 ± 1.48%, while original Masked IRL had the highest strict zero-shot mean, 66.26 ± 1.70%. Both mask-based core methods exceeded LC-RL on this benchmark. The ranking changed with endpoint; there is no universal winner. These are within-group preference-ordering accuracies, not robot task-success rates.')
p('Downstream mask corruption exposed a substantive negative finding: original Masked IRL deteriorated sharply at high error rates. An exploratory uncertainty-aware regularizer used association evidence from existing training demonstrations to soften presumed irrelevance. At 30% bit flips, its adapted mean was 65.47 ± 0.42%, compared with 51.18 ± 0.75% for original Masked IRL. It lost accuracy in the clean condition and has not been confirmed on an independent benchmark. Reduced invariance penalty mass remains a competing explanation.')
p('The submission contains 60 completed V3 learners, all final and validation predictions, selected and adapted checkpoints, frozen configurations, loss traces, dataset hashes and independent audits. The next scientific priority is a fresh confirmatory benchmark with mass-matched controls and more seeds. Original simulation assets would still be required for an exact paper reproduction.')
br();h('Contents',toc_entry=False);toc_anchor=doc.add_paragraph('TOC_INSERT')
br();h('1 Assignment objective')
p('The research objective was to turn a published robotics learning method into an inspectable local experiment: identify its assumptions, reconstruct the computation, evaluate comparable baselines and document what the available evidence supports. A successful reproduction must explain both its performance and the distance between the implemented experiment and the original claim.')
p('The practical outputs are a runnable reward-learning pipeline, explicitly separated adaptation and zero-shot endpoints, machine-readable multi-seed results and an auditable record of problems. The exploratory extension addresses a limitation observed in completed baseline stress tests; it is not used to redefine the primary paper-faithful result.')
h('2 Paper overview')
p('Inverse reinforcement learning infers a reward or cost from demonstrations. Many rewards can explain the same finite trajectories, so demonstration agreement alone does not identify which state variables matter. Language can constrain relevance, but a language-conditioned network may still exploit spurious state correlations.')
p('Masked IRL uses a state-relevance mask to penalize changes in learned reward when irrelevant state coordinates are perturbed. The network retains the full state at inference. Explicit Mask instead removes irrelevant coordinates from the network input. LC-RL conditions on language without either masking mechanism. Oracle masks derive relevance from the known synthetic preference semantics; author-saved LLM masks approximate that information from the committed artifact.')
eq('ℒ = ℒIRL + λ ℒmask       with c = −r')
eq('ℒmask = (1/B) Σᵢ Σₜ Σⱼ (1 − mᵢⱼ) Eε |c(sᵢₜ + εeⱼ, ℓᵢ) − c(sᵢₜ, ℓᵢ)|')
p('Here B counts demonstrations, t indexes waypoints, j indexes state coordinates, m is relevance and ε follows Uniform(0,1) in V3. The expectation is estimated with one draw per demonstration, time and coordinate. Absolute differences occur before summing over time and coordinates. A cost sign reversal leaves this invariance term unchanged. The full IRL estimator is discussed in Section 8.')
br();h('3 Original repository audit')
table('What the released snapshot made reproducible',['Available','Unavailable or incomplete'],[['Reward learners and baseline implementations','Main simulation trajectory arrays and shortest-path arrays'],['Configuration files and preference splits','Simulation resource / URDF package needed for the original scenes'],['Processed real-robot arrays and preprocessing artifacts','Released trained model checkpoints'],['Committed saved semantic masks','Some exact training and estimation protocol details']],[3.45,3.45])
p('The repository does contain data-related artifacts, including processed robot arrays. The missing element is the central simulation data and resource set needed for the requested simulation result. Available real-robot arrays do not substitute for the missing simulation trajectories, and the presence of a data directory is not evidence that every paper experiment is runnable.')
p('The accepted source audit compared all 155 original archive files and confirmed the historical freeze. Its source-line references are retained in evidence/first_checkpoint_source_audit.md. The portable package includes the upstream modules needed by the delivered learners and the associated configuration artifacts; unrelated robot execution scripts are excluded.')
h('4 Reproduction strategy')
p('V2 first reconstructed a functioning released-code baseline. V3 then changed recoverable paper/code discrepancies while keeping the same fixed simulation_v2 trajectories. Separate endpoints resolved the conflict between paper-style adaptation to held-out preferences and a genuinely unadapted preference test. The extension tier followed completion of the primary multi-seed and mask-corruption comparisons.')
table('Completed stages and their role',['Stage','Role','Interpretive boundary'],[['V2','Released-code reconstruction, one seed','Historical context; protocol differs from V3'],['V3 core','Three methods × three seeds','Primary paper-faithful settings on replacement data'],['Strict zero-shot','30 additional preference identities','No demonstrations, gradients or checkpoint selection'],['Stress tests and extensions','Corruption, saved masks, lambda, soft masks','Exploratory benchmark reuse; primary results frozen']],[1.15,2.65,3.1])
br();h('5 Controlled replacement simulation')
p('The replacement simulation is a fixed controlled dataset, not a reconstruction of the authors’ unavailable scene assets. Candidate trajectories were generated in the earlier reconstruction stage, checked for validity and collision constraints, grouped by shared start and goal, and split by scene. This delivery reuses those files without regenerating any trajectories.')
table('Dataset accounting',['Quantity','Verified value'],[['Candidates / rejected / retained','3000 / 784 / 2216'],['Usable scenes','29; one of the original 30 candidate scenes retained no valid trajectories'],['Train / validation / test trajectories','1459 / 359 / 398'],['Train / validation / test scenes','19 / 5 / 5'],['Start-goal groups','491 total: 323 / 81 / 87'],['Waypoints and representation','21 waypoints; 121 raw values; 19 model-state coordinates'],['Semantic features','Five preference features; train-only min/max scaling'],['Endpoint audit','Zero mismatch within each retained start-goal group']],[3.1,3.8])
p('Shared endpoints make pairwise comparison meaningful: alternatives solve the same start-goal problem rather than benefiting from different task requirements. The independent verifier compares the endpoint-related raw coordinates within every group. Trajectory, scene and group identities are all disjoint across splits. Features are scaled using training extrema only; validation and test features are transformed with those frozen values.')
p('Each preference is a five-dimensional nonzero ternary vector. Ground-truth trajectory cost is the dot product of that preference and the five scaled features. Demonstrations are sampled from the training pool with probabilities proportional to exp(−20 × cost), with ten samples per preference. This synthetic preference model makes evaluation inspectable, but also constrains what the benchmark can establish about human behavior.')
p('There are 657 validation and 748 test candidate pairs before preference-dependent true-cost ties are removed. All candidates within a retained group contribute to pair construction. The package includes the consolidated trajectory, state, feature, scene and group arrays once, with the original split and scaling files. Redundant per-scene candidate arrays are excluded and recorded in the packaging policy.')
br();h('6 Released-code V2')
p('V2 established end-to-end feasibility using one seed, 100 epochs, learning rate 0.0001, batch size 64 and trainable T5. It did not use the V3 adaptation stage. Checkpoints were selected at epochs 60 for LC-RL and Masked IRL and 80 for Explicit Mask. V2 validation used the same preference identities later evaluated on different test scenes; V3 replaces that selection design.')
table('Historical V2 ordering accuracy in percent',['Method','Seen','Unseen'],[['LC-RL','68.63','68.14'],['Oracle Masked IRL','71.37','70.35'],['Explicit Mask','71.98','72.83']],[3.5,1.7,1.7])
p('These single-seed values are historical context. They cannot be interpreted as a paired V2-to-V3 performance change because the training, adaptation and selection protocols differ. The original V2 freeze remains unchanged outside this compact V3 delivery.')
h('7 Paper versus released code')
table('Consequential implementation choices',['Component','Manuscript','Released behavior','V3'],[['T5 backbone','Frozen','Trainable','Frozen'],['Mask lambda','10','1','10'],['Perturbation','Uniform','Gaussian','Uniform(0,1)'],['Mask objective','Local state/coordinate','Joint trajectory','Local absolute differences'],['Pretrain + adapt','1000 + 100','Mismatched schedule','1000 + 100'],['Pretraining LR','0.001','Different setting','0.001'],['Fine-tuning','100 epochs','Released mismatch','100; LR 0.0001'],['Importance estimator','Under-specified','Cross-sample ratios','Retained explicitly']],[1.5,1.4,2,2])
p('Freezing changes which representation can adapt. Noise distribution and aggregation change the regularizer’s behavior and scale. Schedules and learning rates change optimization exposure. These are substantive experimental choices, not interchangeable implementation details. The adaptation LR follows the released factor of 0.1; it is not presented as a separately specified manuscript constant.')
br();h('8 V3 implementation')
h('Frozen language features and trainable reward model',2)
p('T5-base and its tokenizer are pinned to revision a9723ea7f1b39c1eae772870f3b547bf6ef7e6c1. T5 runs in evaluation mode with requires_grad=False, is excluded from both optimizers and is checked for absent gradients at each update. Hashes of all backbone state tensors match before and after every completed run. The added language projection, FiLM layers and reward head remain trainable.')
p('The imported FiLMRewardModel uses 19 state coordinates and 128-dimensional conditioning, hidden widths 128/256/128, an initial ReLU followed by Tanh activations and float64 reward weights. The projection maps 768-dimensional pooled language features to 128 dimensions in float32. The upstream reward-model source is unchanged.')
p('Released mean pooling includes padded positions. Frozen features are therefore cached using both exact instruction text and the original dynamically padded batch length. Caching only the instruction would change the representation. A direct-versus-cached comparison runs for each learner, with recorded numerical tolerance. A later CPU lookup memoized token counts; all 203 checked batches retained identical token IDs and attention masks. Both executed source versions remain archived.')
h('Local perturbations and update accounting',2)
p('For each irrelevant coordinate, V3 adds one Uniform(0,1) draw to that coordinate alone. It computes absolute state-cost differences before summing over the 21 waypoints and 19 coordinates, then divides by demonstration batch size only. It neither averages over time/coordinates nor takes an absolute difference after a trajectory sum. This prevents temporal cancellation in the mask term.')
p('Chunk size 4096 controls memory. Each chunk contributes gradients to the same optimizer step, with unchanged loss normalization. Numerical tests compare values and gradients against an unchunked calculation and check perturbation bounds, coordinate isolation and nonmutation. A full 512-example profile completed with finite gradients, about 1.11 GB allocated and 1.50 GB reserved GPU memory.')
p('The configured batch size is 512, but complete demonstration pools contain 400 pretraining and 300 adaptation examples. Each epoch therefore makes one optimizer update on the actual pool. The contrast pool uses 512 samples before adaptation and 300 during adaptation, retaining the released fine-tune conditioning behavior. Training does not pad the demonstration count to 512.')
br();h('The retained estimator and selection boundary',2)
p('The released implementation forms a cross-sample importance expression rather than an elementwise ratio. V3 preserves this behavior and evaluates it in log space to avoid overflow. For a cost vector c with detached p = softmax(−c), its term A(c) is the logarithm of the mean over every pair i,j of exp(−cᵢ)/pⱼ.')
eq('A(c) = logmeanexp(−c) + logmeanexp(−log p)')
eq('ℒIRL = mean(cdemo) + log(exp(A(ccontrast)) + exp(A(cdemo)))')
p('This factorization is algebraically equivalent to the released cross-sample mean. A numerical test verifies both value and gradient equality with the original expression. Detaching the sampling weights preserves the released gradient semantics. It does not establish that the estimator is the uniquely correct interpretation of the manuscript; that ambiguity is a principal reason for the partially paper-faithful label.')
p('Pretraining uses Adam at 0.001 for 1000 epochs. Validation is performed every ten epochs using only the 40 training preference identities in validation scenes. The selected checkpoint maximizes mean ordering accuracy; the first maximum wins exact ties. The final epoch-1000 checkpoint is not automatically selected.')
p('Adaptation starts from the selected pretraining checkpoint, resets Adam and runs exactly 100 epochs at 0.0001 using ten training-pool demonstrations for each of 30 adapted preference identities. The fixed epoch-100 adapted checkpoint is evaluated without adaptation-specific checkpoint selection. All fitting and selection finish before the final test bundle opens.')
p('The final test lock records configuration and selected/adapted checkpoint hashes and a timestamp. The strict zero-shot endpoint uses the selected pretraining weights on a separate set of 30 identities. Neither their demonstrations nor their costs enter optimization or checkpoint choice. “Before adaptation” is an endpoint evaluated from saved weights after fitting finishes; it does not mean test performance was inspected to choose how to adapt.')
table('Implementation evidence to inspect',['Question','Evidence'],[['Was T5 frozen?','Per-run t5_verification.json; before/after tensor hashes and optimizer exclusion'],['Which code ran?','Frozen config source_hashes and script_hashes; matching source_archive files'],['Which weights were tested?','selection.json and FINAL_TEST_STARTED.json'],['Did the schedule complete?','1100 loss rows; runtime optimizer_steps and examples_seen'],['Were all selection scores checked?','100 saved validation NPZs per run; independent verifier']],[2.15,4.75])
br();h('9 Experimental protocol')
table('Primary and extension protocol',['Item','Specification'],[['Seeds','12345, 23456, 34567; fixed in advance'],['Preference roles','40 training; 30 adapted; 30 distinct strict zero-shot'],['Demonstrations','10 per training/adapted identity; 400 + 300 total'],['Schedule','1000 pretraining + 100 adaptation epochs'],['Validation','Every 10 pretraining epochs; training identities in validation scenes'],['Selection','Highest macro preference ordering accuracy; earliest exact tie'],['Test lock','After all fitting; saved pretrain and fixed adaptation checkpoints'],['Common randomness','Matched initialization/demonstrations; deterministic paired corruption'],['Actual exposure','1100 optimizer updates; 430000 demonstration presentations'],['Replicate unit','Training seed; n = 3, not trajectory pairs']],[2.15,4.75])
p('Evaluation constructs all trajectory pairs within each start-goal group. True-cost differences with magnitude at most 10⁻¹⁰ are excluded for that preference. A predicted-cost tie within the same tolerance receives 0.5 credit; other pairs receive one or zero according to ordering agreement. Scores are averaged within each preference, then equally across preferences. No preference receives extra weight merely because it has more non-tied pairs.')
eq('SE = SD / √3     ;     SD² = Σₖ(xₖ − x̄)² / (3 − 1)')
p('CSV tables retain unrounded values. Displayed means and SEs are percentages across the three seeds; paired effects subtract matched seed scores first, then summarize the three differences. SD describes observed seed spread, while SE describes uncertainty of the seed mean under this limited replication. Neither captures uncertainty across different datasets or environment families, and three seeds provide weak support for broad significance claims.')
p('All conditions use the same fixed dataset. Core clean runs are reused at zero corruption rather than repeated to obtain a favorable score. The completed matrix is 9 clean core, 24 additional corruption, 6 saved-mask, 6 lambda and 15 uncertainty-aware learners. Reused clean baselines are references, not extra independent runs. The 60 frozen configurations and every validation/test prediction are included.')
br();h('10 Main V3 results')
table('Core ordering accuracy, mean ± SE percent across three seeds',['Method','Before adaptation','Few-shot adapted','Strict zero-shot'],[[label]+[stat(m,e) for e in EP] for m,label in [('lcrl','LC-RL'),('explicit','Explicit Mask'),('masked','Original Masked IRL')]],[2,1.63,1.63,1.64])
figure('figure1_primary_comparison.png','Figure 1. Existing validated primary comparison, copied unchanged. Error bars show training-seed SE. The endpoints use different preference roles and/or checkpoint states.')
p('Explicit Mask has the highest adapted core mean, while original Masked IRL has the highest strict zero-shot core mean. Both mask-based methods have higher core means than LC-RL on this replacement benchmark. The comparisons are endpoint-specific: the paired Masked-minus-Explicit effect is −0.97 ± 2.19 percentage points after adaptation and +0.81 ± 3.90 points in strict zero-shot evaluation.')
p('These results support the usefulness of relevance information within this controlled setting. They do not establish universal superiority of implicit masking over explicit input masking, reproduce the authors’ original numerical values, or predict physical-robot success. The modest number of seeds and shared replacement dataset limit generalization.')
br();h('11 Adaptation and strict zero-shot')
p('Paper-style adaptation and strict zero-shot answer different questions. The adapted endpoint measures performance after the model receives ten demonstrations and 100 gradient updates for each held-out adaptation identity. Strict zero-shot measures a different set of 30 identities using only the selected pretraining checkpoint. Calling both simply “unseen” would conceal a consequential difference in learning exposure.')
figure('figure2_adaptation.png','Figure 2. Existing validated adaptation comparison, copied unchanged. Before/after adapted-preference scores form a matched comparison; strict zero-shot uses separate identities and pretraining weights.')
p('All three clean core methods have higher average adapted-preference scores after adaptation than before it. This is a descriptive paired benefit in the specified few-shot protocol. It does not imply that the strict endpoint received an equivalent benefit: that endpoint intentionally remains unadapted. The strict score should not be subtracted from the adapted score as if only training exposure changed, because their preference identities also differ.')
p('Validation is restricted to training preference identities, preventing selection on the strict test identities. The strict set is deterministically selected by SHA256 ranking from the saved-mask preference universe after excluding the original training and adapted identities. The exported manifest records the ranks, oracle masks and saved masks so that this partition can be checked independently.')
br();h('12 Author-saved LLM masks')
p('The condition is named author_saved_llm_mask throughout the canonical registry. Masks come from the committed theta_to_pred_mask_sdim19.json artifact, which contains 242 nonzero ternary preference keys. The five preference coordinates are joined into a stable lookup key and mapped to a 19-dimensional state mask. The exact artifact hash and all actual masks used by each run are retained.')
table('Saved-mask ordering accuracy, mean ± SE percent',['Method','Before adaptation','Adapted','Strict zero-shot'],[[label]+[stat(m,e,source='author_saved_llm_mask') for e in EP] for m,label in [('explicit','Explicit Mask'),('masked','Masked IRL')]],[2,1.63,1.63,1.64])
p('No new API calls were made, and the historical LLM model/version provenance was not recreated. Consequently, this condition evaluates a saved artifact, not a fresh GPT reproduction or a current LLM mask-generation pipeline. Reusing the committed masks also makes the downstream experiment independent of API availability or changing model responses.')
p('Both methods degrade substantially relative to their oracle-mask versions on this benchmark. This provides downstream evidence that mask quality matters to learned reward ordering, beyond standalone mask classification metrics. It is not a controlled causal estimate of a single type of semantic error: the saved masks can contain structured mistakes that differ from the synthetic independent-flip model.')
p('Explicit Mask uses the mask in both fitting and inference, so incorrect relevance can remove useful information at evaluation. Original Masked IRL retains the full state but may learn inappropriate invariances from incorrect irrelevance labels. These mechanisms motivate studying downstream behavior directly. The saved-mask result alone does not identify which mechanism dominates for a particular preference.')
h('Traceability and scope',2)
p('The source mask SHA256 is 598ada8219de103109b47c18ebd7a6618fb7fc5b12b5e3f610a4f511e096d36e. The independent verifier reloads this artifact, checks every mapped mask and reconstructs the preference partitions. Evaluating the uncertainty-aware candidate with these saved masks remains deferred; no such result is claimed in this delivery.')
br();h('13 Robustness to mask errors')
p('Independent mask bits are flipped at p = 0, 0.05, 0.10, 0.20 and 0.30. A SHA256-derived seed for each preference and training seed produces a shared uniform vector; thresholding it gives nested corruptions across severity. Explicit and Masked learners therefore receive matched errors. Corruption affects learning and adaptation, and also Explicit Mask inference. This is a downstream reward-learning experiment, not merely a mask-F1 calculation.')
figure('figure3_mask_robustness.png','Figure 3. Existing validated downstream robustness curves, copied unchanged. Zero corruption reuses the clean core learners. Larger flip probabilities do not guarantee monotonic observed scores.')
table('Severe-corruption adapted accuracy, mean ± SE percent',['Flip probability','Explicit Mask','Original Masked IRL'],[[f'{q:.2f}',stat('explicit',EP[1],q),stat('masked',EP[1],q)] for q in [.2,.3]],[2.3,2.3,2.3])
p('Original Masked IRL is not uniformly more robust. At severe corruption its adapted score approaches the 50% pairwise chance reference and can fall well below Explicit Mask. Losses and gradients remained finite; the poor scores are retained scientific findings, not discarded numerical failures.')
p('Corruption changes both error identity and total invariance penalty mass. False irrelevance can penalize truly informative state changes, while false relevance removes intended regularization. Because these effects co-vary, the curves do not prove a unique causal mechanism. Mass-matched controls are required before attributing the degradation solely to semantic error identity.')
br();h('14 Lambda sensitivity')
table('Lambda ablation, mean ± SE percent',['Lambda','Before adaptation','Adapted','Unadapted strict'],[[str(l)]+[ext('lambda',l,e) if l!=10 else stat('masked',e) for e in EP] for l in [1,3,10]],[1.2,1.9,1.9,1.9])
p('Lambda 10 remains the primary manuscript-supported setting. The additional lambda-1 and lambda-3 conditions each used the same three seeds and clean oracle masks; the lambda-10 comparison reuses the original core Masked runs. No new primary result was substituted after observing the ablation.')
p('Lambda 1 has the highest means among these three clean conditions on this reconstruction. This is evidence that a paper-supported hyperparameter need not maximize performance on replacement data. Its effect depends on feature scaling, trajectory length, the number of penalized coordinates and the aggregation convention. Averaging over additional dimensions would also change effective regularization, even if the numeric lambda stayed fixed.')
p('The ablation illustrates why fidelity and performance tuning must be reported separately. Selecting the best observed lambda and relabeling it as the primary reproduction would obscure the question being answered. Here the primary comparison asks how the recoverable manuscript choices behave; the ablation asks how sensitive that result is to regularization strength in the same benchmark.')
h('What the ablation does not establish',2)
p('The three-value grid is not a global hyperparameter search and does not demonstrate an optimal lambda for other datasets, corruption mechanisms or physical tasks. It also does not settle whether the uncertainty-aware method benefits from informative coordinate weights or simply from less total regularization. A mass-matched scalar-lambda control is therefore a high-priority follow-up.')
p('The extension protocol was frozen before extension outcomes were inspected, but after baseline stress results were known. This timing is recorded rather than treated as independent confirmation. All six added lambda learners, including their loss traces and validation selections, remain in the package alongside the 15 uncertainty-aware learners.')
br();h('15 Proposed uncertainty-aware extension')
p('The severe-corruption finding motivated a conservative relaxation of hard irrelevance. The candidate keeps lambda 10 and the local V3 objective but replaces binary relevance b with a soft relevance m. Association evidence can reduce a penalty imposed by a possibly incorrect zero mask entry. It cannot turn an observed relevant coordinate into an irrelevant one.')
eq('m = b + (1 − b)D')
eq('1 − m = (1 − b)(1 − D)')
p('For each coordinate, each trajectory is represented by its mean state over 21 waypoints. D is the maximum absolute difference between the empirical cumulative distribution of ten existing demonstration means and that of the 1459 training-trajectory means. This Kolmogorov–Smirnov-style distance lies in [0,1]. It is an association score; no calibration as causal relevance is assumed.')
eq('Dⱼ = supₓ |Fdemo,j(x) − Ftrain,j(x)|')
p('Pretraining weights use only that preference’s existing pretraining demonstrations and training states. Adaptation weights use the existing adaptation demonstrations only during adaptation. Strict zero-shot identities receive no demonstration-derived D. No oracle reward, corrected oracle mask, validation/test state or newly sampled demonstration is used to fit the confidence weights. Oracle comparisons of errors are post-hoc diagnostics only.')
p('A zero binary entry becomes D rather than remaining zero, reducing its invariance weight from one to 1−D. A binary one stays one regardless of D, so false relevance is not repaired. Correlated coordinates can show association without causal relevance; marginal distributions may miss joint effects; and ten demonstrations give a noisy empirical distribution. The name uncertainty-aware refers to this heuristic relaxation, not calibrated uncertainty.')
p('The formula, seed list, severity grid and selection rule were fixed before candidate results were inspected. Nevertheless, the method was designed after observing baseline failures and evaluated on the same benchmark. It is a promising exploratory extension, not an independently validated improvement or a new state of the art.')
br();h('Exploratory evidence and competing explanations',2)
figure('figure4_uncertainty_aware.png','Figure 4. Existing validated uncertainty-aware comparison, copied unchanged. The candidate improves high-corruption means while losing accuracy in the clean condition.')
table('Paired adapted differences in percentage points, mean ± SE',['Flip probability','Candidate − Masked','Candidate − Explicit'],[[f'{q:.2f}',diff(q,EP[1],'masked'),diff(q,EP[1],'explicit')] for q in [0,.2,.3]],[2.3,2.3,2.3])
p('At 30% corruption, candidate strict zero-shot accuracy exceeds original Masked IRL by '+diff(.3,EP[2],'masked')+' points and Explicit Mask by '+diff(.3,EP[2],'explicit')+' points. These differences are calculated per matched seed. They describe this completed grid; they do not remove benchmark reuse or establish significance across environments.')
p('The candidate also lowers total invariance penalty mass. Mean retained pretraining mass is about 0.629 of binary-mask mass at p=0 and 0.592 at p=0.30. Reduced regularization may account for part of the benefit. Although post-hoc diagnostics show stronger relaxation on falsely irrelevant entries, causal attribution requires controls that match mass while changing how weights are assigned.')
p('The clean adapted effect is negative against both comparators. This trade-off is retained in the figures and tables. Fresh benchmark confirmation, mass-matched controls and evaluation with author-saved LLM masks are needed before elevating the candidate beyond exploratory status.')
br();h('16 What we learned')
h('Scientific lessons',2)
p('Reproduction depends on the evaluation protocol as much as the architecture. The frozen/trainable encoder choice, loss aggregation and preference exposure each change the question an experiment answers. A system called “Masked IRL” can therefore implement materially different learning behavior unless these details are recorded.')
p('Relevance masks are a downstream bottleneck. Oracle-mask advantages did not imply equally strong saved-mask performance or automatic robustness to synthetic errors. Explicit and implicit masking changed rank with endpoint and corruption severity. Negative results made this limitation visible and provided a concrete motivation for the exploratory extension.')
p('Paper fidelity and benchmark performance are distinct. Lambda 10 preserved the primary manuscript-supported comparison even though lambda 1 produced better clean means here. Adaptation and strict zero-shot generalization likewise require separate labels because demonstrations and gradients differ between their preference sets.')
h('Engineering lessons',2)
p('Data provenance must extend beyond filenames. Scene/group split checks, train-only feature scaling and common endpoint verification establish that pairwise comparisons refer to comparable tasks. Freezing the dataset allowed V3 changes to be studied without simultaneously changing the trajectory distribution.')
p('Memory optimization must preserve mathematics. Chunking was accepted only after value and gradient equivalence checks with unchanged batch normalization. Language caching also required a padded-length-aware key because the released pooling includes pads. The tokenization optimization was tested separately and its executed source versions were archived.')
p('Preserved failures make the run history more defensible. The singleton-mask shape failure and incorrect manifest-path launch were documented rather than erased. Neither explains the poor completed high-corruption results, which passed finite-loss and independent metric checks.')
h('Reproducibility lessons',2)
p('An apparent author-code defect should not be silently repaired inside the primary comparison. The cross-sample estimator was retained, stabilized algebraically and explicitly marked for a later ablation. Hashes, exact registries and final-test locks identify which code and checkpoints produced each score, but they cannot compensate for missing original assets or establish out-of-benchmark validity.')
p('A useful submission exposes the limitations alongside successful execution. It reports all prespecified seeds, the negative clean-condition effect of the candidate, and the unresolved penalty-mass confound. This makes later confirmation possible without requiring a reviewer to trust a favorable narrative.')
br();h('17 Problems encountered and solutions')
table('Problems, resolutions and remaining limits',['Problem','Cause','Resolution','Remaining limitation'],[['Missing simulation assets','Unreleased trajectories/resources','Fixed controlled replacement simulation','Exact original result unavailable'],['Invalid or colliding trajectories','Candidate paths violate checks','Reject 784 of 3000; retain 2216','Replacement generator defines distribution'],['Endpoint drift','Alternatives must share a task','Shared endpoint groups; zero-mismatch audit','Only retained trajectories assessed'],['Preference exposure in selection','V2 validation reused later test identities','V3 validation uses training identities; separate strict set','Adapted identities still intentionally receive gradients'],['Local-objective memory','Coordinate/time expansion','Chunk 4096; equivalent gradients; full-512 profile','Runtime depends on hardware'],['Mask-shape failure','Oracle returned (1,19)','Flatten to (19); keep failed profile','Adapter assumption explicit'],['Launch-path error','Wrong preference manifest location','Locate accepted file; verify hash before launch','Portable paths need setup checks'],['Tokenization overhead','Repeated padded-length lookup','Memoize exact counts; 203-batch equivalence check','Released padded pooling retained'],['Unexpected robustness ranking','Incorrect invariance plus changing penalty mass','Report negative results; explore soft weights','Mechanism unresolved without controls']],[1.35,1.6,2,1.95])
p('These entries summarize documented incidents and methodological corrections. The two retained execution failures occurred before a successful learner result was produced. Completed training did not suffer a numerical failure that would justify deleting poor-performing corruption conditions. The execution ledger distinguishes profile, failed launch and completed learner records.')
br();h('18 Limitations')
p('The replacement simulation limits correspondence to the paper’s original task distribution. Only three training seeds were evaluated, and all conditions share one fixed dataset. The released cross-sample importance estimator remains in place. No physical robot experiment or original robot control pipeline was reproduced.')
p('Independent bit flips are a simplified error model that does not reproduce structured language-model mistakes. Exact historical LLM model/version provenance is unavailable. The candidate uses a heuristic marginal association score and cannot repair false relevance. It was developed after baseline inspection and has no fresh independent benchmark evaluation. Error identity and regularization mass remain confounded.')
p('Verification recomputes saved-prediction metrics and checks trace consistency, data partitions, source hashes and frozen-backbone evidence. It does not rerun all models, certify that saved predictions arose from inference without trusting the recorded execution, or establish causal validity. Checkpoint inference requires the separately downloaded pinned T5 weights and a compatible training environment.')
h('19 Prioritized future work')
table('Next studies in priority order',['Priority','Study','Question resolved'],[['Tier 1','Fresh confirmatory benchmark; five seeds','Does the candidate generalize beyond the reused benchmark?'],['Tier 1','Mass-matched corruption controls','Does informed weighting help beyond reduced regularization?'],['Tier 1','Candidate with author-saved LLM masks','Does the gain extend to committed structured mask errors?'],['Tier 2','Corrected estimator; Gaussian/Uniform; frozen/trainable T5','Which retained choices change learning and robustness?'],['Tier 2','Structured mask errors','Does behavior persist under plausible semantic mistakes?'],['Tier 3','Physical robot, calibrated uncertainty, object-semantic masks, broader tasks','Does the method transfer to richer real-world settings?']],[.8,3.05,3.05])
p('None of these deferred studies was run to prepare this submission. A new benchmark and predeclared mass controls should precede claims of a validated algorithmic improvement; original simulation assets remain necessary to assess exact paper reproduction.')
br();h('20 Conclusion')
p('The project successfully implemented and audited a controlled Masked IRL reproduction pipeline, progressing from a released-code reconstruction to a substantially more paper-faithful V3 comparison. The delivered evidence links fixed trajectories, frozen configurations, code versions, selection decisions, checkpoints and raw predictions across 60 completed learners.')
p('The paper/code audit identified consequential differences in encoder training, regularization, noise and optimization schedule. V3 resolved the recoverable differences while retaining and disclosing the released importance estimator. The resulting scientific classification remains a partially paper-faithful reconstruction on a controlled replacement simulation dataset.')
p('Mask-based core methods exceeded LC-RL on this benchmark, but Explicit Mask and original Masked IRL had different advantages by endpoint. Downstream corruption exposed a strong limitation of hard invariance under mask errors. The uncertainty-aware candidate substantially improved high-corruption means while losing clean-condition accuracy and reducing penalty mass. It is promising exploratory evidence with an explicit confirmation agenda.')
p('The strongest contribution of this reproduction is a transparent account of what ran, what changed, what failed and what the results can support. Exact original numerical reproduction remains unavailable without the missing assets. The professor can inspect the short unchanged PDF first, consult this report for reasoning, and use the packaged verifier to audit all completed scores without retraining.')
h('Sources and evidence',2)
p('[1] Masked IRL: LLM-Guided Reward Disambiguation from Demonstrations and Language. Manuscript v1, arXiv:2511.14565. https://arxiv.org/html/2511.14565v1. Paper trace locations: Section IV-B, Section IV-D, Figure 3 and Appendix B; see the accepted source audit.')
p('[2] MIT-CLEAR-Lab/Masked-IRL. https://github.com/MIT-CLEAR-Lab/Masked-IRL. Snapshot b52bc2e3f1ec5597360a74c2641283f6c984a76b; MIT license.')
p('[3] Canonical completed V3 evidence: results/tables/primary and results/tables/extensions; evidence/run_registry.json; per-run config, protocol, selection and raw predictions. All displayed V3 summary statistics are loaded from these machine-readable tables.')
p('[4] Historical reconstruction and decisions: evidence/first_checkpoint_source_audit.md, implementation_decisions_original.md, extension_protocol.md and improvement_decision.md. The original V2 freeze is preserved separately in the local project.')
br();h('Appendix A Environment')
p('Executed learner environment: Python 3.10.21, NumPy 1.26.4, PyTorch 2.5.0+cu124 and Transformers 4.48.1 on Windows, with two NVIDIA RTX 3080 GPUs. Exact exported dependency records accompany a portable setup guide. The training CLI uses CUDA; the independent saved-prediction verifier needs Python and NumPy only. Hardware/driver differences can affect numerical reproducibility.')
p('T5-base revision: a9723ea7f1b39c1eae772870f3b547bf6ef7e6c1. Encoder weights are not duplicated in this submission; the pinned public model must be cached before training or checkpoint inference. No secret token or paid API is needed. See environment/ENVIRONMENT_README.md.')
h('Appendix B Run registry')
table('Unique completed learner count',['Condition','Count','Relationship'],[['Clean core','9','Three methods × three seeds'],['Additional corruption','24','Two methods × four nonzero levels × three seeds'],['Author-saved LLM masks','6','Two methods × three seeds'],['Lambda 1 and 3','6','Two settings × three seeds; lambda 10 reused'],['Uncertainty-aware','15','Five levels × three seeds, including its own clean condition'],['Total','60','39 primary/comparison + 21 extensions']],[2.8,.7,3.4])
p('All configurations reside beside their runs. evidence/run_registry.json maps condition, seed and relative folder. Category index files reference shared zero-corruption baselines rather than duplicating them. Full original source archives preserve executed script hashes; current portable adapters are identified separately.')
h('Appendix C Exact seeds')
p('The prespecified seeds are 12345, 23456 and 34567. Matching seeds share model initialization and demonstration draws across paired conditions. Corruption uniforms are derived from the seed and preference key using SHA256; they are nested across flip probabilities. The strict preference set uses a separately documented deterministic ranking salt, not a favorable-outcome selection.')
br();h('Appendix D Important hashes')
for name,rel in [('Unchanged professor PDF','report/Masked_IRL_V3_Professor_Report.pdf'),('Preference manifest','data_manifest/preference_manifest.json'),('Committed saved masks','code/upstream/config/data_split_config/theta_to_pred_mask_sdim19.json'),('Fixed model states','code/upstream/reconstruction/datasets/simulation_v2/states.npy')]:
 p(name+':');pp=doc.add_paragraph(hashlib.sha256((P/rel).read_bytes()).hexdigest());pp.runs[0].font.name='Consolas';pp.runs[0].font.size=Pt(8)
p('SHA256SUMS.txt covers payload files and DELIVERY_MANIFEST.json. The manifest excludes itself and the checksum list to avoid recursive hashing. The archive SHA256 is supplied in an external sidecar. Hashes establish identity and integrity, not scientific correctness.')
h('Appendix E Reproduction commands')
p('From the extracted package root, install the environment as described in environment/ENVIRONMENT_README.md. These commands are portable entry points; training commands are provided for future use and were not run during packaging.')
for line in ['python scripts/check_environment.py','python verification/verify_submission.py --package .','python scripts/run_core_v3.py --output work/core --gpus 0','python scripts/run_corruption.py --output work/core --gpus 0','python scripts/run_extensions.py --base-root work/core --output work/extensions --gpus 0','python scripts/summarize_results.py --package . --output work/recomputed_tables']:
 pp=doc.add_paragraph(line);pp.runs[0].font.name='Consolas';pp.runs[0].font.size=Pt(8)
p('Training outputs must go to a new work directory. The launcher uses existing trainer and scheduler interfaces and leaves packaged evidence unchanged. Saved-result evaluation is included in the independent verifier; checkpoint inference has a separate evaluation entry point documented in README.md.')
h('Appendix F Package structure')
p('START_HERE.md gives the reading order. report/ and figures/ hold the documents and unchanged figures. code/ contains upstream dependencies, portable adapters and the fixed dataset. data_manifest/ indexes the data; results/ holds 60 runs and canonical tables; verification/ supplies independent audits; evidence/ records decisions; environment/ and scripts/ provide setup and execution commands.')
p('Environment binaries, model caches, redundant epoch-1000 checkpoints and duplicate candidate arrays are excluded. All 6000 validation files, 300 final prediction files and 120 selected/adapted checkpoints are retained for independent inspection.')

# Deterministic static linked TOC, page numbers verified against the rendered artifact.
anchor=toc_anchor._p
for title,n,key in toc:
 pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(3);pp.paragraph_format.line_spacing=1
 pp.paragraph_format.tab_stops.add_tab_stop(Inches(6.6),WD_ALIGN_PARAGRAPH.RIGHT)
 link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),key);r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=title;r.append(t);link.append(r);pp._p.append(link);pp.add_run('\t'+str(n));anchor.addprevious(pp._p)
anchor.getparent().remove(anchor)
a.output.parent.mkdir(parents=True,exist_ok=True);doc.save(a.output)
(a.output.parent/'toc_expected.json').write_text(json.dumps(toc,indent=2),encoding='utf-8')
print(json.dumps({'docx':str(a.output),'planned_pages':page,'tables':table_no,'toc_entries':len(toc)}))
